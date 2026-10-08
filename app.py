import io
import json
import os

import streamlit as st
from docx import Document
from google import genai
from google.genai import types
from pypdf import PdfReader
from pydantic import BaseModel, Field


st.set_page_config(
    page_title="AI Resume ATS Checker",
    page_icon="📄",
    layout="wide",
)

MODEL_NAME = "gemini-2.5-flash"


class ResumeAnalysis(BaseModel):
    ats_score: int = Field(description="Overall ATS-readiness score from 0 to 100.")
    summary: str = Field(description="Short overall assessment of the resume.")
    strengths: list[str] = Field(description="3 to 6 strong points in the resume.")
    improvements: list[str] = Field(description="5 to 8 specific improvements, written as actionable advice.")
    missing_keywords: list[str] = Field(description="Important missing keywords. Use an empty list when no job description is supplied and no clear keywords can be inferred.")
    formatting_issues: list[str] = Field(description="Potential ATS-unfriendly formatting or structure issues.")
    section_feedback: dict[str, str] = Field(
        description="Feedback for Summary/Profile, Experience, Education, Skills, Projects, and Certifications. Use 'Not provided' where a section is absent."
    )
    ats_breakdown: dict[str, int] = Field(
        description="Scores from 0 to 100 for readability, keyword_alignment, section_structure, skills, experience, and education."
    )


def get_api_key() -> str | None:
    """Read the Gemini key from Streamlit Secrets or an environment variable."""
    try:
        key = st.secrets.get("GEMINI_API_KEY")
    except Exception:
        key = None

    return key or os.getenv("GEMINI_API_KEY")


def extract_text(uploaded_file) -> str:
    """Extract text from PDF, DOCX, or TXT uploads."""
    file_bytes = uploaded_file.getvalue()
    filename = uploaded_file.name.lower()

    if filename.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(file_bytes))
        pages = []
        for page in reader.pages:
            pages.append(page.extract_text() or "")
        return "\n".join(pages)

    if filename.endswith(".docx"):
        document = Document(io.BytesIO(file_bytes))
        paragraphs = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                paragraphs.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(paragraphs)

    if filename.endswith(".txt"):
        return file_bytes.decode("utf-8", errors="replace")

    raise ValueError("Unsupported file type. Please upload a PDF, DOCX, or TXT file.")


def clean_text(text: str) -> str:
    """Normalize extracted resume text and keep the prompt reasonably sized."""
    lines = [line.strip() for line in text.splitlines()]
    text = "\n".join(line for line in lines if line)
    return text[:60000]


def analyze_resume(resume_text: str, job_description: str) -> ResumeAnalysis:
    api_key = get_api_key()
    if not api_key:
        raise RuntimeError(
            "Gemini API key not found. Add GEMINI_API_KEY to Streamlit Secrets "
            "or set it as an environment variable."
        )

    client = genai.Client(api_key=api_key)

    job_context = (
        f"\nTARGET JOB DESCRIPTION:\n{job_description[:20000]}"
        if job_description.strip()
        else "\nNo job description was supplied. Evaluate general ATS-readiness."
    )

    prompt = f"""
You are an expert ATS resume reviewer and career coach.

Analyze the resume below for ATS-readiness. The ATS score is an estimate, not a score
from any particular company's proprietary ATS. Do not invent facts about the candidate.

Use these principles:
- A strong ATS resume should be machine-readable, clearly structured, and easy to parse.
- Check standard sections, headings, dates, job titles, skills, measurable achievements,
  consistency, and keyword alignment.
- If a target job description is supplied, compare the resume against it and identify
  useful missing keywords. Do not recommend keyword stuffing or false claims.
- Do not penalize a candidate simply for not having information that is not expected.
- Give practical improvements that the candidate can actually make.
- Keep feedback concise and specific.
- Scores must be integers from 0 to 100.

RESUME:
{resume_text}

{job_context}
"""

    response = client.models.generate_content(
        model=MODEL_NAME,
        contents=prompt,
        config=types.GenerateContentConfig(
            response_mime_type="application/json",
            response_schema=ResumeAnalysis,
            temperature=0.2,
        ),
    )

    if not response.text:
        raise RuntimeError("Gemini returned an empty response.")

    return ResumeAnalysis.model_validate_json(response.text)


def score_label(score: int) -> str:
    if score >= 85:
        return "Excellent ATS readiness"
    if score >= 70:
        return "Good ATS readiness"
    if score >= 55:
        return "Needs improvement"
    return "Major improvements recommended"


st.title("📄 AI Resume ATS Checker")
st.write(
    "Upload your resume to get an estimated ATS score, strengths, formatting issues, "
    "missing keywords, and practical improvement suggestions."
)

with st.sidebar:
    st.header("Settings")
    st.info(
        "Your Gemini API key is read from Streamlit Secrets. "
        "Do not put the key directly inside app.py or commit it to GitHub."
    )
    st.caption(f"Model: {MODEL_NAME}")

uploaded_file = st.file_uploader(
    "Upload your resume",
    type=["pdf", "docx", "txt"],
    help="Supported formats: PDF, DOCX and TXT.",
)

job_description = st.text_area(
    "Optional: paste the target job description",
    height=220,
    placeholder="Adding the job description makes keyword alignment more useful.",
)

if uploaded_file:
    try:
        resume_text = clean_text(extract_text(uploaded_file))

        if len(resume_text) < 100:
            st.warning(
                "Very little text could be extracted. If this is a scanned/image-only PDF, "
                "use a text-based PDF or DOCX."
            )
        else:
            with st.expander("Preview extracted resume text"):
                st.text(resume_text[:8000])

            if st.button("🔍 Analyze Resume", type="primary", use_container_width=True):
                with st.spinner("Analyzing your resume with Gemini..."):
                    try:
                        result = analyze_resume(resume_text, job_description)

                        st.session_state["analysis"] = result
                        st.session_state["filename"] = uploaded_file.name
                    except Exception as exc:
                        st.error(f"Analysis failed: {exc}")

    except Exception as exc:
        st.error(f"Could not read this file: {exc}")

result = st.session_state.get("analysis")

if result:
    st.divider()
    st.subheader(f"Results for {st.session_state.get('filename', 'resume')}")

    col1, col2 = st.columns([1, 2])
    with col1:
        st.metric("Estimated ATS Score", f"{result.ats_score}/100")
    with col2:
        st.success(score_label(result.ats_score))
        st.write(result.summary)

    st.subheader("📊 ATS Breakdown")
    breakdown = result.ats_breakdown
    cols = st.columns(3)
    items = [
        ("Readability", "readability"),
        ("Keyword Alignment", "keyword_alignment"),
        ("Section Structure", "section_structure"),
        ("Skills", "skills"),
        ("Experience", "experience"),
        ("Education", "education"),
    ]
    for i, (label, key) in enumerate(items):
        with cols[i % 3]:
            value = int(max(0, min(100, breakdown.get(key, 0))))
            st.progress(value / 100)
            st.caption(f"{label}: {value}/100")

    left, right = st.columns(2)

    with left:
        st.subheader("✅ Strengths")
        for item in result.strengths:
            st.markdown(f"- {item}")

        st.subheader("⚠️ Formatting / ATS Issues")
        if result.formatting_issues:
            for item in result.formatting_issues:
                st.markdown(f"- {item}")
        else:
            st.write("No major formatting issues were identified.")

    with right:
        st.subheader("🚀 Improvements")
        for item in result.improvements:
            st.markdown(f"- {item}")

        st.subheader("🔑 Missing Keywords")
        if result.missing_keywords:
            st.write(", ".join(result.missing_keywords))
        else:
            st.write("No specific missing keywords were identified.")

    st.subheader("🧩 Section Feedback")
    for section, feedback in result.section_feedback.items():
        st.markdown(f"**{section}:** {feedback}")

    st.caption(
        "Note: This is an AI-generated ATS-readiness estimate. Real ATS systems differ, "
        "and a score cannot guarantee that an application will pass screening."
    )
else:
    st.info("Upload a resume and click “Analyze Resume” to see your results.")
